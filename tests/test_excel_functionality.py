"""
EXCEL FUNCTIONALITY TESTS - ensures all positive functionalities are
passing
"""

import openpyxl
import pytest

from expenses import (
    open_file,
    get_or_create_sheet,
    write_headers,
    get_sheet_name,
    save_expense_to_excel

)

#1.test to check all the created headers are correct first is Date then User,Category,Amount,Note
@pytest.mark.excel_functionality
def test_headers_are_correct(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"

    save_expense_to_excel("daniel", "car", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    assert ws["A1"].value == "Date"
    assert ws["B1"].value == "User"
    assert ws["C1"].value == "Category"
    assert ws["D1"].value == "Amount"
    assert ws["E1"].value == "Note"

#2.test to check that data is inserted into correct columns "daniel" to User etc ....
@pytest.mark.excel_functionality
def test_data_is_written(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"
    save_expense_to_excel("daniel", "car", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    assert ws["A2"].value == "22 April 2026"
    assert ws["B2"].value == "daniel"
    assert ws["C2"].value == "car"
    assert ws["D2"].value == 50
    assert ws["E2"].value == "gas"

#3.test to check that we can save multiple   expenses and they will be inserted to correct raw
@pytest.mark.excel_functionality
def test_multiple_expenses_append_correctly(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"
    save_expense_to_excel("daniel", "car", 50, "gas", str(file)) #first expense
    save_expense_to_excel("inbar", "groceries", 50, "coffee", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    """First expense expectations"""
    assert ws["A2"].value == "22 April 2026"
    assert ws["B2"].value == "daniel"
    assert ws["C2"].value == "car"
    assert ws["D2"].value == 50
    assert ws["E2"].value == "gas"
    """Second expense expectations"""
    assert ws["A3"].value == "22 April 2026"
    assert ws["B3"].value == "inbar"
    assert ws["C3"].value == "groceries"
    assert ws["D3"].value == 50
    assert ws["E3"].value == "coffee"