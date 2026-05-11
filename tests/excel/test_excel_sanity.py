"""SANITY TESTS TO CHECK CORE FUNCTIONALITY OF
   EXPENSES EXCEL FUNCTIONS IF DONT PASS TESTING SHOULD STOP TILL FIXED !"""

import openpyxl
import pytest

from expenses_excel import (
    open_file,
    get_or_create_sheet,
    write_headers,
    get_sheet_name,
    save_expense_to_excel

)

#1.test to check open_file() creates a file and it exists
@pytest.mark.excel_sanity
def test_file_is_created(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date) # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name) # mock sheet name  "April 2026"

    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))
    assert file.exists()


#2. test to check sheet is created eg "April 2026"
@pytest.mark.excel_sanity
def test_sheet_is_created(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date) # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)   # mock sheet name  "April 2026"

    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    assert "April 2026" in wb.sheetnames

#3.test to check headers created at all
@pytest.mark.excel_sanity
def test_headers_are_written(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"

    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    assert ws["A1"].value is not None

#4.test to check that some data is inserted at all into the excel and is not None
@pytest.mark.excel_sanity
def test_data_is_written(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"
    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    assert ws["A2"].value is not None #check there is something inside the first raw first column

