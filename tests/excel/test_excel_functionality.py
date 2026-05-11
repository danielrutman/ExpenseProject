"""
EXCEL FUNCTIONALITY TESTS - ensures all positive functionalities are
passing
"""

import openpyxl
import pytest

from expenses_excel import (
    open_file,
    get_or_create_sheet,
    write_headers,
    get_sheet_name,
    save_expense_to_excel

)

#1.test to check all the created headers are correct first is Date then User,Category,Amount,Note
@pytest.mark.excel_functionality
def test_headers_are_correct(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"

    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    assert ws["A1"].value == "Date"
    assert ws["B1"].value == "User"
    assert ws["C1"].value == "Category"
    assert ws["D1"].value == "Amount"
    assert ws["E1"].value == "Note"

#2.test to check that data is inserted into correct columns "דניאל" to User etc ....
@pytest.mark.excel_functionality
def test_data_is_written(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"
    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    assert ws["A2"].value == "22 April 2026"
    assert ws["B2"].value == "דניאל"
    assert ws["C2"].value == "רכב"
    assert ws["D2"].value == 50
    assert ws["E2"].value == "gas"

#3.test to check that we can save multiple   expenses and they will be inserted to correct raw
@pytest.mark.excel_functionality
def test_multiple_expenses_append_correctly(tmp_path, monkeypatch, mock_date, mock_sheet_name):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"
    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file)) #first expense
    save_expense_to_excel("ענבר", "מצרכים ופארם", 50.5, "coffee", str(file))
    wb = openpyxl.load_workbook(str(file))
    ws = wb["April 2026"]
    """First expense expectations"""
    assert ws["A2"].value == "22 April 2026"
    assert ws["B2"].value == "דניאל"
    assert ws["C2"].value == "רכב"
    assert ws["D2"].value == 50
    assert ws["E2"].value == "gas"
    """Second expense expectations"""
    assert ws["A3"].value == "22 April 2026"
    assert ws["B3"].value == "ענבר"
    assert ws["C3"].value == "מצרכים ופארם"
    assert ws["D3"].value == 50.5
    assert ws["E3"].value == "coffee"

#4.  test to check new month creates new spreadsheet
@pytest.mark.excel_functionality
def test_new_month_sheet(tmp_path, monkeypatch, mock_date, mock_sheet_name,mock_new_month_sheet_name,mock_new_month_date):
    file = tmp_path / "test.xlsx"  # creates path but not file yet

    #first  mock expense in april
    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_date)  # mock date "22 April 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_sheet_name)  # mock sheet name  "April 2026"
    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))  # first expense
    #second mock expense in May
    monkeypatch.setattr("expenses_excel.get_current_date", lambda: mock_new_month_date)  # mock date "22 May 2026"
    monkeypatch.setattr("expenses_excel.get_sheet_name", lambda: mock_new_month_sheet_name)  # mock sheet name  "May 2026"
    save_expense_to_excel("דניאל", "רכב", 50, "gas", str(file))  # second expense

    wb = openpyxl.load_workbook(str(file))
    assert "April 2026" in wb.sheetnames
    assert "May 2026" in wb.sheetnames