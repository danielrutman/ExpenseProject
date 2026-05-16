""" EXCEL EXPENSES FUNCTION FILE contains
all the  EXCEL important functions for expenses project"""

import openpyxl
import os

from expenses_functions.expenses import get_current_date
from config import EXCEL_HEADER, MONTHLY_BUDGET





def open_file(FILE_PATH):
    """Opens existing Excel file or creates new one if doesn't exist.
    Returns workbook object."""

    if os.path.exists(FILE_PATH):
        return openpyxl.load_workbook(FILE_PATH)
    else:
        wb = openpyxl.Workbook()
        del wb[wb.sheetnames[0]]  # delete default sheet otherwise will have 2 sheets from start
        return wb

def get_or_create_sheet(wb, sheet_name):
    """Gets existing sheet or creates new one if doesn't exist.
    Returns worksheet object."""

    # wb.sheetnames  returns a list of all worksheet names in a loaded workbook
    if sheet_name in wb.sheetnames:
        return wb[ sheet_name]
    else:
        # wb.create_sheet(sheet_name) method in openpyxl is used to add a new worksheet to an existing workbook object
        return wb.create_sheet(sheet_name)


def write_headers(ws):
    """Writes column headers to row 1 of the sheet.
    Only called when sheet is newly created."""

    #excel columns start at index 1 thats why start is 1 for col, header in  will insert eg: c1 = "Date",c2 = "User" etc
    for col, header in enumerate(EXCEL_HEADER, start=1):
        #ws.cell For each header in our list, write it into row 1 at the correct column position.
        ws.cell(row=1, column=col, value=header)

def get_sheet_name():
    """Returns current month and year as sheet name.
    eg: 'April 2026'"""

    from datetime import datetime
    return datetime.now().strftime("%B %Y")  # "April 2026"



def save_expense_to_excel(name, category, amount, note, FILE_PATH):
    """saves expense to excel file."""

    sheet_name = get_sheet_name()

    #1. open or create excel file
    wb = open_file(FILE_PATH)
    #2. check if sheet exists create it if doesnt eg: april 2026
    ws = get_or_create_sheet(wb, sheet_name)
    if ws["A1"].value is None:
        #3. create headers if the sheet is new from const list EXCEL_HEADER
        write_headers(ws)
    #4.append data to sheet new raw
    ws.append([get_current_date(), name, category, amount, note])
    wb.save(FILE_PATH)
    # 5. calculate remaining budget for current month
    # Amount is in column 4 (D), skip header row 1
    total_spent = sum(
        ws.cell(row=row, column=4).value
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=4).value is not None
    )
    remaining = MONTHLY_BUDGET - total_spent
    return remaining

def generate_report(FILE_PATH):
    """returns monthly spending report per category + total spending"""

    sheet_name = get_sheet_name()
    wb = open_file(FILE_PATH)

    if sheet_name not in wb.sheetnames: # check to see sheet exists ex april 2026
        return "📊 No expenses found for this month yet!"

    ws = wb[sheet_name]

    # sum amounts per category — col 3 = category, col 4 = amount
    category_totals = {}
    for row in range(2, ws.max_row + 1):
        cat = ws.cell(row=row, column=3).value # grab category valuge ex רכב
        amt = ws.cell(row=row, column=4).value # grab amount value for this category column 4 header == amount
        if cat and amt: # will return false if one of the values is empty
            category_totals[cat] = category_totals.get(cat, 0) + amt

    total_spent = sum(category_totals.values()) # returns the total spending till now
    category_lines = "\n".join(f"• {cat} — {amt} NIS" for cat, amt in category_totals.items())

    return (
        f"📊 {sheet_name}\n"
        f"💸 Spent: {total_spent} / {MONTHLY_BUDGET} NIS\n\n"
        f"{category_lines}"
    )
